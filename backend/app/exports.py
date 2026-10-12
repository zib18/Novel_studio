"""Authenticated exports of current or explicitly approved revisions."""
from io import BytesIO
from typing import Literal
from urllib.parse import quote

from docx import Document
from docx.shared import Pt
from fastapi import APIRouter, Depends
from fastapi.responses import Response
from sqlalchemy import and_, select
from sqlalchemy.orm import Session

from app.auth import get_current_user
from app.database import get_db
from app.models import Chapter, Revision
from app.services import owned_novel

router = APIRouter(prefix="/api/v1", tags=["export"])


@router.get("/novels/{novel_id}/export")
def export_novel(
    novel_id: str,
    format: Literal["txt", "docx"] = "txt",
    selection: Literal["current", "approved"] = "current",
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    novel = owned_novel(db, novel_id, user)
    revision_column = Chapter.current_revision_id if selection == "current" else Chapter.approved_revision_id
    rows = db.execute(
        select(Chapter.number, Chapter.title, Revision.text)
        .join(Revision, and_(Revision.id == revision_column, Revision.chapter_id == Chapter.id))
        .where(Chapter.novel_id == novel.id)
        .order_by(Chapter.number, Chapter.id)
    ).all()
    texts = [(f"Глава {number}. {title}", text) for number, title, text in rows]
    if format == "txt":
        content = (novel.title + "\n\n" + "\n\n".join(title + "\n\n" + text for title, text in texts)).encode("utf-8-sig")
        media_type = "text/plain; charset=utf-8"
    else:
        document = Document()
        document.styles["Normal"].font.name = "Times New Roman"
        document.styles["Normal"].font.size = Pt(12)
        document.add_heading(novel.title, level=0)
        for title, text in texts:
            document.add_heading(title, level=1)
            for paragraph in text.split("\n"):
                document.add_paragraph(paragraph)
        buffer = BytesIO()
        document.save(buffer)
        content = buffer.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    filename = "".join(c for c in novel.title if c not in '\r\n/\\').strip()[:100] or "Novel"
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename=novel.{format}; filename*=UTF-8''{quote(filename, safe='')}.{format}"},
    )
